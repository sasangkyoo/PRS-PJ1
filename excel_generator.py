import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
import os
from typing import List, Dict

def _apply_excel_formatting(file_path: str, header_color: str):
    """
    Applies formatting to the Excel file.
    header_color: ARGB hex string (e.g., 'FFDDDDDD' for grey, 'FFADD8E6' for light blue)
    """
    wb = load_workbook(file_path)
    ws = wb.active
    
    # Header formatting
    header_fill = PatternFill(start_color=header_color, end_color=header_color, fill_type="solid")
    header_font = Font(bold=True)
    alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = alignment
        
    # Data rows formatting and auto-adjust column width
    data_alignment = Alignment(vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = data_alignment
            
    # Simple column width adjustment
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min((max_length + 2), 50) # Cap at 50 to prevent excessively wide columns
        ws.column_dimensions[column].width = adjusted_width
        
    wb.save(file_path)

def generate_function_definition_excel(data: List[Dict], output_path: str):
    """
    Generates the Function Definition Excel file.
    """
    # Define exact columns as per spec
    columns = [
        "슬라이드 번호", "화면명", "기능 식별자", "화면 영역", "기능명", 
        "발생 조건", "사전 조건", "정상 처리 결과", "유효성 검증 및 예외 사항"
    ]
    
    # Flatten the data structure if it came directly from JSON
    flat_data = []
    for item in data:
        # Assuming the JSON has {"슬라이드_번호": X, "화면명": Y, "기능_목록": [{...}]}
        slide_num = item.get("슬라이드_번호", "")
        screen_name = item.get("화면명", "")
        
        for func in item.get("기능_목록", []):
            flat_item = {
                "슬라이드 번호": slide_num,
                "화면명": screen_name,
                "기능 식별자": func.get("기능_식별자", ""),
                "화면 영역": func.get("화면_영역", ""),
                "기능명": func.get("기능명", ""),
                "발생 조건": func.get("발생_조건", ""),
                "사전 조건": func.get("사전_조건", ""),
                "정상 처리 결과": func.get("정상_처리_결과", ""),
                "유효성 검증 및 예외 사항": func.get("유효성_검증_및_예외_로직", "")
            }
            flat_data.append(flat_item)
            
    df = pd.DataFrame(flat_data)
    for col in columns:
        if col not in df.columns:
            df[col] = ""
    df = df[columns]
    
    df.to_excel(output_path, index=False)
    
    # Apply formatting (Grey background for header)
    _apply_excel_formatting(output_path, header_color="FFD9D9D9")
    
def generate_test_case_excel(data: Dict, output_path: str):
    """
    Generates the Test Case Excel file.
    """
    # Define exact columns as per spec
    columns = [
        "슬라이드 번호", "화면명", "테스트 식별자", "연관 기능 식별자", "테스트 유형", 
        "테스트 목적", "사전 조건", "테스트 절차", "화면의 기대 결과"
    ]
    
    # Data is expected to be a dict with "테스트_케이스_목록"
    test_cases = data.get("테스트_케이스_목록", [])
    
    flat_data = []
    for item in test_cases:
        flat_item = {
            "슬라이드 번호": item.get("슬라이드_번호", ""),
            "화면명": item.get("화면명", ""),
            "테스트 식별자": item.get("테스트_식별자", ""),
            "연관 기능 식별자": item.get("참조_기능_식별자", ""),
            "테스트 유형": item.get("테스트_유형", ""),
            "테스트 목적": item.get("테스트_목적", ""),
            "사전 조건": item.get("사전_조건", ""),
            "테스트 절차": item.get("테스트_절차", ""),
            "화면의 기대 결과": item.get("기대_결과", "")
        }
        flat_data.append(flat_item)
            
    df = pd.DataFrame(flat_data)
    for col in columns:
        if col not in df.columns:
            df[col] = ""
    df = df[columns]
    
    df.to_excel(output_path, index=False)
    
    # Apply formatting (Blue background for header)
    _apply_excel_formatting(output_path, header_color="FFB4C6E7")
